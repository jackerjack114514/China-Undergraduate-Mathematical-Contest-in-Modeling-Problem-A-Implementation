import numpy as np
import openpyxl
import pandas as pd
import os
import time
from pathlib import Path

"""依旧温度与水分耦合，药材中心即为含水率最高处，取其含水率不高于0.15kg/kg的时刻即可。
P.S.但是否必须严格小于？"""

相对路径=Path(__file__).resolve().parent #必须改成相对路径！
sr=str(相对路径/"附件"/"附件1.xlsx")
sc=str(相对路径/"result3.xlsx")


R=0.02  #药材半径/m
L=0.25  #药材长度/m

#dt、N、pc关键 POINT  稳定性与误差分析需要调整这三个参数！！
dt=5.0                #时间步长/s  #问题1：dt=0.05 问题2：0.125s  #必须取60的约数，否则报错！
N=800                 #径向网格区间数（不能取400以下）
dr=R/N                #径向步长/m
tm=7*24*3600.0        #最长计算时间（实际达不到）
nt=round(tm/dt)       #最大时间步数
ot=round(60.0/dt)
oh=round(21600.0/dt)
nr=20                 #结果径向区间数（0.001m/格）

h=25.0      #表面热对流系数/[W/(m²·K)]
hm=8e-7     #表面水扩散系数/(m/s)
pc=1e-8     #Picard收敛值ε   #1e^10在T较大时无法有效收敛
pm=100      #最大Picard迭代数（实测20~50不够用，迭代成本太高，是否存在均衡点？）
Cl=0.15     #含水率阈值/(kg/kg)，到达后结束

sj=pd.read_excel(sr).iloc[:,:3].to_numpy(dtype=float)
sj=sj[np.isfinite(sj).all(axis=1)]
sj=sj[np.argsort(sj[:,0])]
t=sj[:,0]    #时间/s
Ta=sj[:,1]   #空气温度/℃
Ca=sj[:,2]   #空气水分浓度/(kg/kg)

r=np.arange(N+1,dtype=float)*dr  #径向节点/m
ii=np.arange(1,N,dtype=float)    #内部节点编号
gl=(1-1/(2*ii))/dr**2            #径向向左差分
gr=(1+1/(2*ii))/dr**2            #径向向右差分

def mat(T,C):
    Cs=np.maximum(C,1e-12)    #必须取非0的小值，不然除以0报错!!
    Tk=T+273.15               #开氏度...
    rho=650+128*C                      #药材密度/(kg/m³)
    cp=1450+2736*C/(C+1)               #比热容/[J/(kg·K)]
    k=0.21+0.38*C/(C+1)                #导热系数/[W/(m·K)]
    D=2.4e-3*np.exp(-0.45/Cs-3850/Tk)  #水分扩散系数/(m²/s)
    return rho,cp,k,D

def sanduijiao(dl,d,du,b):  #dl、d、du：下、主、上对角；b：右端向量
    n=len(d)
    dl=dl.copy();d=d.copy();du=du.copy()  #三对角
    b=b.copy()

    for i in range(1,n):
        m=dl[i-1]/d[i-1]
        d[i]-=m*du[i-1]
        b[i]-=m*b[i-1]
    x=np.empty(n,dtype=float)  #解向量
    x[-1]=b[-1]/d[-1]
    for i in range(n-2,-1,-1):
        x[i]=(b[i]-du[i]*x[i+1])/d[i]
    return x

def fd(u,a,S,be,ua): 
    af=(a[:-1]+a[1:])/2      #半节点
    al=dt*gl*af[:-1]/S[1:N]  #向左
    ar=dt*gr*af[1:]/S[1:N]   #向右

    dl=np.zeros(N)   #下对角
    d=np.zeros(N)    #主对角
    du=np.zeros(N)   #上对角
    b=u[:-1].copy()  #右端向量

    a0=dt*4*af[0]/(S[0]*dr**2)  #圆心差分系数
    d[0]=1+a0
    du[0]=-a0

    d[1:N-1]=1+al[:-1]+ar[:-1]  #内部节点平铺
    dl[:N-1]=-al
    du[1:N-1]=-ar[:-1]
    d[-1]=1+al[-1]+ar[-1]
    dl[N-2]=-al[-1]

    v=3+2*dr*be
    p=4/v
    q=-1/v
    s=2*dr*be*ua/v

    d[-1]+=-ar[-1]*p
    dl[N-2]+=-ar[-1]*q
    b[-1]+=ar[-1]*s

    x=sanduijiao(dl,d,du,b)  #内部节点解
    return np.append(x,p*x[-1]+q*x[-2]+s)

#-----------------------------------------------到此为止均沿用问题2的微分方程


time_start=time.time()  #显示计算耗时用，可删去
T0=28.0                 #初始温度/℃
C0=2.55                 #初始含水率/(kg/kg)
T=np.full(N+1,T0)
C=np.full(N+1,C0)
S1=np.ones(N+1)         #水分方程的储能系数

to=[0.0]
Co=[C[::N//nr].copy()]         #含水率输出数组
pis=0                          #Picard迭代次数和
pim=0                          #最大Picard迭代次数

Ta0=50.00036     #延拓温度--NEW
Ca0=0.0499913    #延拓湿度--NEW

for it in range(1,nt+1):  #时间推进
    tn=it*dt              #当前新时刻/s
    Cp=C.copy()           #上一时刻的含水率
    Ti=T.copy()           #温度Picard初值
    Ci=C.copy()           #含水率Picard初值

    '''Tai=float(np.interp(tn,t,Ta))  #当前空气温度/℃     #弃用
    Cai=float(np.interp(tn,t,Ca))  #当前空气水分浓度/(kg/kg)'''

    Tai=float(np.interp(tn,t,Ta,right=Ta0))   #取均值延拓，具体数值由excel给出
    Cai=float(np.interp(tn,t,Ca,right=Ca0))

    for ip in range(1,pm+1):
        rhoi,cpi,ki,Di=mat(Ti,Ci)   
        Si=rhoi*cpi                       #体积热容/[J/(m³·K)]
        Tn=fd(T,ki,Si,h/ki[-1],Tai)

        rhoi,cpi,ki,Di=mat(Tn,Ci) 
        Cn=fd(C,Di,S1,hm/Di[-1],Cai)

        er=max(np.max(np.abs(Tn-Ti))/50,np.max(np.abs(Cn-Ci))/C0) 
        Ti=Tn  #温度
        Ci=Cn  #含水率

        if er<pc:
            break
    else:
        raise RuntimeError("t={:.6f}s时Picard迭代在{}次内未收敛。".format(tn,pm))  #简单break仍会执行后续输出，必须ERROR
    
    #尽量避免因迭代次数超标而break的情况发生，这一问取上限pm尽可能大，合理性可严格证明。

    T=Ti              #温度已收敛
    C=Ci              #含水率已收敛
    pis+=ip
    pim=max(pim,ip)   #最大Picard迭代次数

    if np.max(C)<=Cl:
        mk=Cp>Cl
        if np.any(mk):
            th=np.max((Cp[mk]-Cl)/(Cp[mk]-C[mk])) 
            th=float(np.clip(th,0,1))
        else:
            th=0.0
        te=(it-1+th)*dt          #结束时间/s
        Ce=Cp+th*(C-Cp)          #结束时刻的含水率
        if abs(te-to[-1])<1e-10:
            Co[-1]=Ce[::N//nr].copy()
        else:
            to.append(te)
            Co.append(Ce[::N//nr].copy())
        break

    if it%ot==0:
        to.append(tn)
        Co.append(C[::N//nr].copy())

    if it%oh==0:
        print("已计算{:.0f}h，中心含水率={:.6f}".format(tn/3600,C[0]))
else:
    raise RuntimeError("计算7天后药材各处含水率仍未降至0.15kg/kg以下。")

to=np.asarray(to,dtype=float)
Co=np.asarray(Co,dtype=float)  #含水率

def xlsx(nm,z,tt):
    sh=ans.create_sheet(nm)
    nf="0.0000"  #所有结果必须保留四位小数（防止28.0000输出成28）
    sh.cell(1,1,"t(s)/r(cm)")
    for j in range(nr+1):
        sh.cell(1,j+2,float(j/10))      #不加这项就会把28.0000输出成28
    for i in range(len(tt)):
        c=sh.cell(i+2,1,float(i*60))       #头一行与头一列不需要4位浮点  #i*60：秒数；i：分钟数
        for j in range(nr+1):
            c=sh.cell(i+2,j+2,round(float(z[i,j]),4))
            c.number_format=nf

ans=openpyxl.Workbook()
w0=ans.active
ans.remove(w0)
xlsx("水分浓度",Co,to)
ans.save(sc)

ri=[0,5,10,15,20]
print("\n表5：药材烘干过程的水分浓度/(kg/kg)")
print("时间/h       0cm       0.5cm       1cm       1.5cm       2cm")
for hs in range(6,int(te//21600)*6+1,6):
    ir=int(hs*3600/60)
    print("{:6.1f}  ".format(float(hs))+"  ".join("{:.4f}".format(x) for x in Co[ir,ri]))
print("结束{:6.4f}  ".format(te/3600)+"  ".join("{:.4f}".format(x) for x in Co[-1,ri]))

time_end=time.time()
time0=time_end-time_start
print("\n计算完成，总耗时{:.4f}秒".format(time0))
print("烘干结束时间：{:.4f}h（{:.4f}s）".format(te/3600,te))
print("烘干结束时中心含水率：{:.10f}kg/kg".format(np.max(Ce)))
print("Picard平均迭代次数：{:.4f}".format(pis/it))
print("Picard最大迭代次数：{}".format(pim))
print("结果已保存至：{}".format(sc))
